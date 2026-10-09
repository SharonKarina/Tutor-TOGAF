from __future__ import annotations

import hashlib
import re
import unicodedata
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".html", ".htm"}

# pypdf does not extract labels from this vector diagram. Keep a page-attributed
# transcription so retrieval can answer questions about the ADM cycle order.
PDF_FIGURE_TRANSCRIPTIONS = {
    ("TOGAF_V9.1-Introduccion.pdf", 4): (
        "Figure 5-1 Architecture Development Cycle. The diagram shows the Preliminary "
        "Phase before the ADM cycle. The lettered cycle proceeds through Phase A: "
        "Architecture Vision; Phase B: Business Architecture; Phase C: Information "
        "Systems Architectures; Phase D: Technology Architecture; Phase E: Opportunities "
        "and Solutions; Phase F: Migration Planning; Phase G: Implementation Governance; "
        "and Phase H: Architecture Change Management. Requirements Management is shown "
        "at the center of the cycle with connections to the phases; it is not a lettered "
        "phase. The supplied Phase C extracts cover Data Architecture and Application "
        "Architecture."
    ),
}


@dataclass(frozen=True)
class SourceDocument:
    text: str
    source: str
    page: int | None = None


@dataclass(frozen=True)
class DocumentChunk:
    chunk_id: str
    text: str
    metadata: dict[str, str | int]


def _clean_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).replace("\x00", " ")
    text = re.sub(r"(?<=\w)-\s*\n\s*(?=\w)", "", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]+", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_documents(data: bytes, filename: str) -> list[SourceDocument]:
    """Extract page-aware text from PDF, DOCX, TXT, or HTML bytes."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Formato no compatible: {suffix or '(sin extensión)'}")

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        source_name = Path(filename).name
        pages = []
        for index, page in enumerate(reader.pages, start=1):
            # TOGAF PDFs position glyphs individually; pypdf's default extraction
            # can split ordinary words ("contin uous", "go vernance"). Layout mode
            # uses page coordinates to restore much more reliable word boundaries.
            text = _clean_text(page.extract_text(extraction_mode="layout") or "")
            figure_transcription = PDF_FIGURE_TRANSCRIPTIONS.get((source_name, index))
            if figure_transcription:
                # Put the short transcription first so chunking keeps all figure
                # labels together in the first chunk for this cited page.
                text = _clean_text(f"{figure_transcription}\n\n{text}")
            pages.append(SourceDocument(text=text, source=source_name, page=index))
        documents = [page for page in pages if page.text]
    elif suffix == ".docx":
        from docx import Document

        document = Document(BytesIO(data))
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            blocks.extend(
                " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                for row in table.rows
            )
        text = _clean_text("\n\n".join(blocks))
        documents = [SourceDocument(text=text, source=Path(filename).name)] if text else []
    elif suffix in {".html", ".htm"}:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(data, "html.parser")
        for element in soup(["script", "style", "noscript", "svg"]):
            element.decompose()
        text = _clean_text(soup.get_text("\n", strip=True))
        documents = [SourceDocument(text=text, source=Path(filename).name)] if text else []
    else:
        try:
            raw_text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raw_text = data.decode("cp1252", errors="replace")
        text = _clean_text(raw_text)
        documents = [SourceDocument(text=text, source=Path(filename).name)] if text else []

    if not documents:
        raise ValueError(f"No se pudo extraer texto de {Path(filename).name}.")
    return documents


def load_directory(directory: str | Path) -> list[SourceDocument]:
    """Read supported files from a directory in deterministic name order."""
    root = Path(directory)
    if not root.exists():
        raise FileNotFoundError(f"No existe la carpeta de documentos: {root}")

    documents: list[SourceDocument] = []
    files = sorted(
        path for path in root.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
    )
    if not files:
        raise FileNotFoundError(
            f"No hay PDF, DOCX, TXT ni HTML en {root}. Copia allí los documentos del corpus."
        )

    for path in files:
        documents.extend(extract_documents(path.read_bytes(), path.name))
    return documents


def _chunk_end(text: str, start: int, target: int) -> int:
    """Prefer paragraph, sentence, or word boundaries near the target size."""
    if target >= len(text):
        return len(text)
    lower_bound = start + max(1, (target - start) // 2)
    boundaries = [
        text.rfind("\n\n", lower_bound, target),
        text.rfind(". ", lower_bound, target),
        text.rfind("? ", lower_bound, target),
        text.rfind("! ", lower_bound, target),
        text.rfind("; ", lower_bound, target),
        text.rfind(" ", lower_bound, target),
    ]
    boundary = max(boundaries)
    if boundary < lower_bound:
        return target
    return boundary + 1 if text[boundary:boundary + 1] in ".?!;" else boundary


def chunk_documents(
    documents: list[SourceDocument], chunk_size: int = 900, overlap: int = 120
) -> list[DocumentChunk]:
    """Split each source page separately with a small character overlap."""
    if chunk_size < 100:
        raise ValueError("chunk_size debe ser de al menos 100 caracteres.")
    if overlap < 0 or overlap >= chunk_size:
        raise ValueError("overlap debe ser mayor o igual a cero y menor que chunk_size.")

    chunks: list[DocumentChunk] = []
    for document in documents:
        text = _clean_text(document.text)
        start = 0
        local_index = 0
        while start < len(text):
            end = _chunk_end(text, start, min(start + chunk_size, len(text)))
            excerpt = text[start:end].strip()
            if excerpt:
                metadata: dict[str, str | int] = {
                    "source": Path(document.source).name,
                    "chunk_index": local_index,
                }
                if document.page is not None:
                    metadata["page"] = document.page
                digest = hashlib.sha256(
                    f"{document.source}|{document.page}|{local_index}|{excerpt}".encode("utf-8")
                ).hexdigest()[:20]
                chunks.append(DocumentChunk(chunk_id=digest, text=excerpt, metadata=metadata))
                local_index += 1
            if end >= len(text):
                break
            next_start = max(start + 1, end - overlap)
            while next_start < end and text[next_start].isspace():
                next_start += 1
            start = next_start
    return chunks
